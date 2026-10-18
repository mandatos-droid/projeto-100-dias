"""Gera os dados do mapa de indicador social (pergunta m1.q8) a partir da planilha do IDSC-BR.

Uso:  python scripts/build_social.py ["../Dados_IDSC-BR_2026 (1).xlsx"]
Saída em public/data/social/:
  catalogo.json        ODS e indicadores (nome, fonte, ano, sentido)
  idsc.json            nota geral do IDSC-BR 2026
  ods-<n>.json         nota de cada ODS (1 a 17)
  <codigo>.json        nota normalizada (0-100, maior = melhor) de cada indicador, por município e por UF
  perfil/<uf>.json     perfil de cada município: população, IDSC, nível, posição e as 17 notas de ODS
A nota de cada UF é a média das notas dos municípios ponderada pela população 2022: é um cálculo da
plataforma, não um número oficial do IDSC.
"""
import json, os, sys
from collections import Counter
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)
XLSX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(RAIZ, '..', 'Dados_IDSC-BR_2026 (1).xlsx')
SAIDA = os.path.join(RAIZ, 'public', 'data', 'social')

ODS = {1: 'Erradicação da pobreza', 2: 'Fome zero e agricultura sustentável', 3: 'Saúde e bem-estar', 4: 'Educação de qualidade',
       5: 'Igualdade de gênero', 6: 'Água potável e saneamento', 7: 'Energia limpa e acessível', 8: 'Trabalho decente e crescimento econômico',
       9: 'Indústria, inovação e infraestrutura', 10: 'Redução das desigualdades', 11: 'Cidades e comunidades sustentáveis',
       12: 'Consumo e produção responsáveis', 13: 'Ação contra a mudança do clima', 14: 'Vida na água', 15: 'Vida terrestre',
       16: 'Paz, justiça e instituições eficazes', 17: 'Parcerias e meios de implementação'}

def num(x):
    try:
        return None if x is None or x == '' else float(x)
    except (TypeError, ValueError):
        return None

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
livro = list(wb.worksheets[0].iter_rows(values_only=True))
cab = [str(c) for c in livro[0]]
ix = {k: cab.index(k) for k in ('ODS', 'Arquivo', 'Indicador', 'Escala', 'FONTE')}
indic = [r for r in livro[1:] if r[ix['Arquivo']]]

linhas = wb.worksheets[1].iter_rows(values_only=True)
h = [str(c) for c in next(linhas)]
col = {c: i for i, c in enumerate(h)}
dados = [r for r in linhas if r[0]]
print(len(dados), 'municípios;', len(indic), 'indicadores')

def val_ind(cod):
    c = col.get(f'Normalizado (0-100): {cod}')
    return c

def grava(nome, obj):
    caminho = os.path.join(SAIDA, nome)
    os.makedirs(os.path.dirname(caminho), exist_ok=True)
    with open(caminho, 'w', encoding='utf-8') as f:
        json.dump(obj, f, ensure_ascii=False, separators=(',', ':'))

def serie(c, ano=None, unidade='nota 0–100'):
    mun, soma, pesos = {}, {}, {}
    for r in dados:
        v = num(r[c])
        if v is None:
            continue
        cod, uf, pop = str(r[0]), r[2], num(r[3]) or 0
        mun[cod] = round(v, 1)
        soma[uf] = soma.get(uf, 0) + v * pop
        pesos[uf] = pesos.get(uf, 0) + pop
    out = {'unidade': unidade, 'mun': mun, 'uf': {u: round(soma[u] / pesos[u], 1) for u in soma if pesos[u]}}
    if ano:
        out['ano'] = ano
    return out

cat = {'fonte': 'IDSC-BR 2026 · Instituto Cidades Sustentáveis', 'url': 'https://idsc.cidadessustentaveis.org.br/',
       'ods': [{'n': n, 'rot': r} for n, r in ODS.items()], 'indicadores': []}

grava('idsc.json', serie(col['IDSC-BR_2026']))
for n in ODS:
    grava(f'ods-{n}.json', serie(col[f'ODS {n}']))

for r in indic:
    cod = str(r[ix['Arquivo']])
    c = val_ind(cod)
    if c is None:
        print('sem coluna:', cod)
        continue
    anos = Counter(str(r2[col[f'year_{cod}']]) for r2 in dados if r2[col[f'year_{cod}']] is not None)
    ano = anos.most_common(1)[0][0] if anos else None
    grava(f'{cod}.json', serie(c, ano))
    cat['indicadores'].append({'id': cod, 'ods': int(float(r[ix['ODS']])), 'nome': str(r[ix['Indicador']]).strip(),
                               'fonte': str(r[ix['FONTE']] or '').strip(), 'ano': ano,
                               'sentido': 'menos é melhor' if str(r[ix['Escala']]).startswith('desc') else 'mais é melhor'})
grava('catalogo.json', cat)

# perfil por UF
niv, pos = col['Nível de Desenvolvimento Sustentável'], col['Classificação 2026']
por_uf = {}
for r in dados:
    por_uf.setdefault(r[2], {})[str(r[0])] = {
        'n': r[1], 'p': int(num(r[3]) or 0), 'i': round(num(r[col['IDSC-BR_2026']]) or 0, 1), 'v': r[niv], 'r': int(num(r[pos]) or 0),
        'o': [None if num(r[col[f'ODS {n}']]) is None else round(num(r[col[f'ODS {n}']]), 1) for n in ODS]}
for uf, d in por_uf.items():
    grava(f'perfil/{uf.lower()}.json', d)
print('ok:', len(cat['indicadores']), 'indicadores,', len(por_uf), 'UFs')
